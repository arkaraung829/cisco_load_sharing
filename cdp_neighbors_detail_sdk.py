#!/usr/bin/env python3
"""
SDK-based variant of cdp_neighbors_detail.py.

Collects 'show cdp neighbors detail' from a list of switches via Catalyst
Center's Command Runner API, using the official catalystcentersdk
(CatalystCenterAPI) instead of raw requests calls.

*** THIS IS A SIDE-BY-SIDE EXPERIMENT, NOT A REPLACEMENT. ***
Run this against the same device(s) as cdp_neighbors_detail.py and diff the
two output CSVs before trusting this version for anything real. The SDK
method/parameter names below are a best-effort mapping of the Command
Runner / Task / File REST endpoints to catalystcentersdk calls - they have
NOT been verified against your installed SDK version. Your own
dnac_onboarding.py already hit a case where an SDK method's parameters
didn't match the REST API docs (discovery.start_discovery had to be
bypassed with a raw requests call) - so this script assumes the same kind
of mismatch is possible here too, and is built to fail loudly and
specifically instead of silently.

Before trusting the output, run:
    python cdp_neighbors_detail_sdk.py --inspect --host ... --username ... --password ...
This prints the REAL method signatures your installed SDK provides for
device lookup, Command Runner, Task, and File - compare them against what
this script calls (see resolve_method() candidate lists below) and tell me
if anything doesn't line up so I can fix the exact call.

Output columns match the raw-requests version exactly:
    Source Device, Neighbor ID, Neighbor IP, Platform, Hostname
Written to cdp_neighbors_results_sdk.csv (different filename from the raw
version's output, so running both side by side doesn't clobber either).

Requires: catalystcentersdk  (pip install catalystcentersdk)
"""

import argparse
import csv
import getpass
import inspect
import json
import os
import re
import socket
import sys
import time

try:
    from catalystcentersdk import CatalystCenterAPI
except ImportError:
    sys.exit("ERROR: catalystcentersdk is not installed - run: pip install catalystcentersdk")

COL_IP = "ip address"
COMMAND = "show cdp neighbors detail"
TASK_POLL_INTERVAL = 5
TASK_POLL_TIMEOUT = 300
DNS_TIMEOUT = 3
BATCH_SIZE = 20    # Command Runner's documented limit: max 5 commands x 20 devices per request
BATCH_PAUSE = 2    # seconds between batches
OUTPUT_FILE = "cdp_neighbors_results_sdk.csv"


def load_dotenv():
    """Load KEY=VALUE pairs from a .env (or credential.env) file next to this
    script into os.environ. Real environment variables take priority."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    env_path = next((p for p in (os.path.join(script_dir, name) for name in (".env", "credential.env"))
                     if os.path.isfile(p)), None)
    if env_path is None:
        return
    with open(env_path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip().strip("'\"")
            if key and key not in os.environ:
                os.environ[key] = value


def parse_args():
    load_dotenv()
    p = argparse.ArgumentParser(description="SDK-based CDP neighbor collection (experimental - see docstring).")
    p.add_argument("--host", default=os.environ.get("DNAC_HOST"),
                   help="Catalyst Center hostname or IP, no scheme (or set DNAC_HOST)")
    p.add_argument("--file", help="Input .xlsx or .csv file with an 'IP Address' column (not needed with --inspect)")
    p.add_argument("--username", default=os.environ.get("DNAC_USER"), help="API username (or set DNAC_USER)")
    p.add_argument("--password", default=os.environ.get("DNAC_PASSWORD"), help="API password (or set DNAC_PASSWORD)")
    p.add_argument("--insecure", "-k", action="store_true", help="Skip TLS certificate verification (self-signed labs)")
    p.add_argument("--no-dns", action="store_true", help="Skip reverse-DNS hostname resolution (faster)")
    p.add_argument("--inspect", action="store_true",
                   help="Print the real SDK method signatures for the calls this script uses, then exit")
    args = p.parse_args()

    if not args.host:
        args.host = input("Catalyst Center host/IP: ")
    if not args.username:
        args.username = input("Catalyst Center username: ")
    if not args.password:
        args.password = getpass.getpass("Catalyst Center password: ")
    if not args.inspect and not args.file:
        sys.exit("ERROR: --file is required unless --inspect is used")
    return args


def normalize_host(host):
    """Strip any scheme the caller already included (DNAC_HOST=https://... is
    a common credential.env mistake) so we don't end up with https://https://."""
    return re.sub(r"^\s*https?://", "", host.strip(), flags=re.IGNORECASE).rstrip("/")


def build_client(host, username, password, verify):
    return CatalystCenterAPI(
        base_url=f"https://{normalize_host(host)}",
        username=username,
        password=password,
        verify=verify,
        wait_on_rate_limit=True,   # SDK handles 429 backoff itself - no hand-rolled retry needed
    )


def unwrap(response):
    """SDK responses are typically Box-like objects with a .response attribute;
    some calls (e.g. File API) return the raw content directly instead."""
    if hasattr(response, "response"):
        return response.response
    if isinstance(response, dict):
        return response.get("response", response)
    return response


def resolve_method(obj, candidates, label):
    """Find the first attribute name that actually exists on the installed
    SDK object, instead of assuming a single guessed name is correct."""
    for name in candidates:
        if hasattr(obj, name):
            return getattr(obj, name), name
    available = [m for m in dir(obj) if not m.startswith("_")]
    sys.exit(
        f"ERROR: none of {candidates} exist under api.{label} in your installed catalystcentersdk.\n"
        f"Available methods on api.{label}: {available}\n"
        f"Update the candidate list in resolve_method() for '{label}' to the correct name."
    )


def call_sdk(method, label, **kwargs):
    """Call an SDK method; on a parameter mismatch, print its real signature
    instead of a bare traceback, so a fix is one glance away."""
    try:
        return method(**kwargs)
    except TypeError as exc:
        sig = inspect.signature(method)
        sys.exit(
            f"ERROR: {label} rejected the parameters {list(kwargs.keys())}.\n"
            f"  {exc}\n"
            f"  Actual signature: {label}{sig}\n"
            f"Fix the kwarg names passed to this call to match, or run --inspect for the full picture."
        )


def inspect_sdk(api):
    print("Real SDK method signatures for the calls this script uses:\n")
    for label, obj, candidates in (
        ("devices", api.devices, ["get_device_list", "get_device_by_ip", "get_network_device_by_ip"]),
        ("command_runner", api.command_runner,
         ["run_read_only_commands_on_devices_to_get_their_real_time_information",
          "run_read_only_commands_on_devices"]),
        ("task", api.task, ["get_task_by_id"]),
        ("file", api.file, ["download_a_file_by_fileid", "download_a_file_by_file_id", "get_file"]),
    ):
        method, name = resolve_method(obj, candidates, label)
        print(f"api.{label}.{name}{inspect.signature(method)}")
    print("\nCompare these against the kwargs used in this script's get_device_id(), "
          "run_command(), wait_for_task(), and get_file() functions.")


DEBUG = os.environ.get("CDP_SDK_DEBUG") == "1"


def debug(msg):
    if DEBUG:
        print(f"[debug] {msg}")


# -- Step 3: IP -> device UUID (via SDK) -------------------------------------
def get_device_id(api, ip):
    method, name = resolve_method(api.devices, ["get_device_list", "get_device_by_ip", "get_network_device_by_ip"], "devices")
    debug(f"devices call resolved to '{name}'")
    response = call_sdk(method, f"devices.{name}", managementIpAddress=ip)
    result = unwrap(response)
    if isinstance(result, list):
        result = result[0] if result else None
    if not result:
        return None, None
    return result.get("id"), result.get("hostname", ip)


# -- Step 4: run the CLI command on all devices in one call -----------------
def run_command(api, device_ids, command):
    method, name = resolve_method(
        api.command_runner,
        ["run_read_only_commands_on_devices_to_get_their_real_time_information",
         "run_read_only_commands_on_devices"],
        "command_runner",
    )
    debug(f"command_runner call resolved to '{name}'")
    response = call_sdk(method, f"command_runner.{name}", commands=[command], deviceUuids=device_ids)
    task = unwrap(response)
    task_id = task.get("taskId") if isinstance(task, dict) else getattr(task, "taskId", None)
    if not task_id:
        # Per-batch problem, not a setup problem - raise so the caller can
        # skip just this batch instead of aborting the whole (possibly
        # 80+ batch) run.
        raise RuntimeError(f"no taskId returned from command_runner.{name}: {task}")
    return task_id


# -- Step 5: poll the task until it finishes ---------------------------------
def wait_for_task(api, task_id):
    method, name = resolve_method(api.task, ["get_task_by_id"], "task")
    deadline = time.time() + TASK_POLL_TIMEOUT
    while time.time() < deadline:
        response = call_sdk(method, f"task.{name}", task_id=task_id)
        task = unwrap(response)
        is_error = task.get("isError") if isinstance(task, dict) else getattr(task, "isError", False)
        end_time = task.get("endTime") if isinstance(task, dict) else getattr(task, "endTime", None)
        if is_error:
            reason = task.get("failureReason") if isinstance(task, dict) else getattr(task, "failureReason", task)
            raise RuntimeError(f"Command Runner task failed: {reason}")
        if end_time:
            progress = task.get("progress") if isinstance(task, dict) else getattr(task, "progress", "")
            try:
                file_id = json.loads(progress).get("fileId")
            except (ValueError, TypeError, AttributeError):
                file_id = None
            if not file_id:
                raise RuntimeError(f"task finished but no fileId in progress: {progress}")
            return file_id
        time.sleep(TASK_POLL_INTERVAL)
    raise RuntimeError(f"Command Runner task {task_id} did not finish within {TASK_POLL_TIMEOUT}s")


def extract_file_content(response):
    """The File API can return a raw download-wrapper object (Cisco's
    DownloadResponse) instead of already-parsed JSON, since file downloads
    can be any content type. Try the standard accessors in order and
    json-decode whichever one yields usable text/bytes."""
    if isinstance(response, (list, dict)):
        return response

    json_method = getattr(response, "json", None)
    if callable(json_method):
        try:
            return json_method()
        except Exception as exc:
            debug(f"response.json() raised: {exc}")

    for attr in ("data", "content", "text", "body", "raw_response"):
        value = getattr(response, attr, None)
        if value is None:
            continue
        if isinstance(value, (list, dict)):
            return value
        if isinstance(value, (bytes, bytearray)):
            try:
                return json.loads(value.decode("utf-8"))
            except Exception as exc:
                debug(f"response.{attr} (bytes) json.loads failed: {exc}")
        elif isinstance(value, str):
            try:
                return json.loads(value)
            except Exception as exc:
                debug(f"response.{attr} (str) json.loads failed: {exc}")

    debug(f"could not extract content from {type(response).__name__}; "
          f"available attrs: {[m for m in dir(response) if not m.startswith('_')]}")
    return []


# -- Step 6: fetch the raw CLI output -----------------------------------------
def get_file(api, file_id):
    method, name = resolve_method(
        api.file, ["download_a_file_by_fileid", "download_a_file_by_file_id", "get_file"], "file")
    debug(f"file call resolved to '{name}'")
    response = call_sdk(method, f"file.{name}", file_id=file_id)
    debug(f"raw response from file.{name}: type={type(response).__name__} "
          f"repr={repr(response)[:1000]}")
    result = unwrap(response)
    result = extract_file_content(result)
    debug(f"after extract_file_content(): type={type(result).__name__} repr={repr(result)[:1000]}")
    return result


# -- Step 7: parse 'show cdp neighbors detail' text into structured rows ----
def parse_cdp_neighbors_detail(text):
    """Split raw CDP output into one dict per neighbor entry."""
    blocks = re.split(r"-{5,}", text)
    entries = []
    for block in blocks:
        if "Device ID:" not in block:
            continue
        device_id = re.search(r"Device ID:\s*(\S+)", block)
        ip_addr = re.search(r"IP address:\s*([\d.]+)", block)
        platform = re.search(r"Platform:\s*(.+?),\s*Capabilities:", block)
        entries.append({
            "neighbor_id": device_id.group(1).strip() if device_id else "",
            "neighbor_ip": ip_addr.group(1).strip() if ip_addr else "",
            "platform": platform.group(1).strip() if platform else "",
        })
    return entries


def reverse_dns(ip):
    if not ip:
        return ""
    try:
        return socket.gethostbyaddr(ip)[0]
    except (socket.herror, socket.gaierror, socket.timeout):
        return ""


def chunked(seq, size):
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


# -- Step 2: read the input file for IP addresses ----------------------------
def read_ips(path):
    """Yield IP addresses from an .xlsx or .csv file with an 'IP Address' column."""
    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("ERROR: reading .xlsx needs openpyxl - run: pip install openpyxl "
                     "(or export the sheet as CSV instead)")
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = ([("" if c is None else str(c).strip()) for c in row]
                for row in ws.iter_rows(values_only=True))
    else:
        rows = ([c.strip() for c in row] for row in csv.reader(open(path, newline="")))

    header = next(rows, None)
    if not header:
        sys.exit(f"ERROR: {path} is empty.")
    lowered = [h.lower() for h in header]
    try:
        ip_col = lowered.index(COL_IP)
    except ValueError:
        sys.exit(f"ERROR: {path} must have an '{COL_IP}' column header (found: {header})")

    seen = set()
    for row in rows:
        if len(row) <= ip_col:
            continue
        ip = row[ip_col]
        if ip and ip not in seen:
            seen.add(ip)
            yield ip


def main():
    args = parse_args()
    socket.setdefaulttimeout(DNS_TIMEOUT)

    print(f"Connecting to https://{normalize_host(args.host)} ...")
    api = build_client(args.host, args.username, args.password, verify=not args.insecure)
    print("Authenticated OK\n")

    if args.inspect:
        inspect_sdk(api)
        return

    ips = list(read_ips(args.file))
    if not ips:
        sys.exit(f"ERROR: no usable rows found in {args.file}")
    print(f"Loaded {len(ips)} device IP(s) from {args.file}\n")

    devices = []   # (ip, device_id, hostname)
    skipped = []
    for ip in ips:
        device_id, hostname = get_device_id(api, ip)
        if device_id is None:
            print(f"!! {ip}: not found in inventory - skipping")
            skipped.append(ip)
            continue
        print(f"ok {ip}: {hostname} -> deviceId {device_id}")
        devices.append((ip, device_id, hostname))
    print()

    if not devices:
        sys.exit("ERROR: no devices resolved to inventory entries - nothing to query.")

    device_ids = [d[1] for d in devices]
    id_to_source = {d[1]: (d[0], d[2]) for d in devices}

    batches = list(chunked(device_ids, BATCH_SIZE))
    print(f"Running '{COMMAND}' on {len(device_ids)} device(s) in {len(batches)} "
          f"batch(es) of up to {BATCH_SIZE} (Command Runner's per-request device limit)...")

    script_dir = os.path.dirname(os.path.abspath(__file__))
    output_path = os.path.join(script_dir, OUTPUT_FILE)
    dns_cache = {}
    row_count = 0
    empty_devices = []
    failed_batches = []

    # Write incrementally, one batch at a time, so a crash/interrupt partway
    # through a large run (e.g. 1770 devices = ~89 batches) still leaves a
    # usable CSV with everything completed so far, instead of losing it all.
    with open(output_path, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Source Device", "Neighbor ID", "Neighbor IP", "Platform", "Hostname"])
        f.flush()

        for i, batch in enumerate(batches, 1):
            print(f"\nBatch {i}/{len(batches)} ({len(batch)} device(s))...")
            try:
                task_id = run_command(api, batch, COMMAND)
                print(f"  taskId: {task_id} - waiting for completion...")
                file_id = wait_for_task(api, task_id)
                batch_results = get_file(api, file_id)
            except SystemExit:
                raise
            except Exception as exc:
                print(f"  !! batch {i} failed, skipping it and continuing: {exc}")
                failed_batches.append(i)
                if i < len(batches):
                    time.sleep(BATCH_PAUSE)
                continue

            if not batch_results:
                print(f"  !! no content returned for this batch (type={type(batch_results).__name__}, "
                      f"value={batch_results!r}). Re-run with CDP_SDK_DEBUG=1 set for full detail.")

            for entry in batch_results or []:
                entry = dict(entry) if not isinstance(entry, dict) else entry
                dev_uuid = entry.get("deviceUuid")
                source_ip, source_hostname = id_to_source.get(dev_uuid, ("?", dev_uuid or "?"))
                command_responses = entry.get("commandResponses") or {}
                raw = (command_responses.get("SUCCESS") or {}).get(COMMAND)
                if not raw:
                    err = (command_responses.get("FAILURE") or {}).get(COMMAND) \
                        or (command_responses.get("BLACKLISTED") or {}).get(COMMAND) \
                        or "no output returned"
                    print(f"  !! {source_hostname} ({source_ip}): {err}")
                    continue

                neighbors = parse_cdp_neighbors_detail(raw)
                if not neighbors:
                    empty_devices.append(source_hostname)
                    continue

                for n in neighbors:
                    hostname = ""
                    if not args.no_dns and n["neighbor_ip"]:
                        if n["neighbor_ip"] not in dns_cache:
                            dns_cache[n["neighbor_ip"]] = reverse_dns(n["neighbor_ip"])
                        hostname = dns_cache[n["neighbor_ip"]]
                    writer.writerow([source_hostname, n["neighbor_id"], n["neighbor_ip"], n["platform"], hostname])
                    row_count += 1
                    print(f"  {source_hostname} -> {n['neighbor_id']} ({n['neighbor_ip']}) "
                          f"[{n['platform']}] hostname={hostname}")
            f.flush()

            if i < len(batches):
                time.sleep(BATCH_PAUSE)

    print("\n" + "=" * 60)
    print(f"Wrote {row_count} neighbor entries to {output_path}")
    if failed_batches:
        print(f"{len(failed_batches)} batch(es) failed entirely and were skipped: {failed_batches}")
    if empty_devices:
        print(f"{len(empty_devices)} device(s) had no CDP neighbors: {', '.join(empty_devices)}")
    if skipped:
        print(f"{len(skipped)} IP(s) skipped (not in inventory): {', '.join(skipped)}")


if __name__ == "__main__":
    main()
