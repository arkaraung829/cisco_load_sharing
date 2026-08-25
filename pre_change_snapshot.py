#!/usr/bin/env python3
"""
Pre-change snapshot: gather standard show commands from a list of switches
via Catalyst Center's Command Runner API (SDK-based, same pattern as
cdp_neighbors_detail_sdk.py) and save one text file per device.

Commands gathered (all read-only):
    show version
    show switch
    show interfaces status
    show process cpu history
    show logging last 200
    show running-config

NOT included: 'write memory'. That persists the running config to the
device's own startup-config - it's a write operation, not a read, and
Command Runner's API is explicitly scoped to READ-ONLY commands (see its
own name: "run read only commands..."). It's very likely rejected by
Command Runner's allowlist by policy, not just by omission. Before
scripting it: test 'write memory' manually in Catalyst Center's Command
Runner GUI tool first (Tools > Command Runner) the same way 'show cdp
neighbors detail' was validated earlier - if it's accepted there, tell
me and I'll add it as a separate, explicitly-confirmed write action
(same treatment as activate_golden_image.py's confirmation gate). If
it's rejected, Catalyst Center's own Configuration Archive feature
(visible on the Activities page as a scheduled weekly job) is the
correct native mechanism for persisting/backing up device config - that
can likely be triggered on demand instead.

Workflow:
    1. Authenticate            (catalystcentersdk)
    2. Read IP list from the input file
    3. Resolve IP -> device
    4. For each command, in batches of up to 20 devices (Command Runner's
       per-request device limit): run it, poll the task, fetch the output
    5. Write one text file per device combining all commands' output, in
       pre_change_snapshots/ next to this script
    6. Print a summary of what succeeded / failed / was skipped

Usage:
    python pre_change_snapshot.py --file devices.csv
    python pre_change_snapshot.py --file devices.csv --insecure

Credentials and host are taken from (highest priority first):
    1. command-line flags --host / --username / --password
    2. environment variables DNAC_HOST / DNAC_USER / DNAC_PASSWORD
    3. a .env (or credential.env) file in the same folder as this script

Requires: catalystcentersdk  (pip install catalystcentersdk)
"""

import argparse
import csv
import datetime
import getpass
import inspect
import json
import os
import re
import sys
import time

try:
    from catalystcentersdk import CatalystCenterAPI
except ImportError:
    sys.exit("ERROR: catalystcentersdk is not installed - run: pip install catalystcentersdk")

COL_IP = "ip address"

COMMANDS = [
    "show version",
    "show switch",
    "show interfaces status",
    "show process cpu history",
    "show logging last 200",
    "show running-config",
]

TASK_POLL_INTERVAL = 5
TASK_POLL_TIMEOUT = 300
BATCH_SIZE = 20    # Command Runner's documented limit: max 5 commands x 20 devices per request
BATCH_PAUSE = 2    # seconds between batches
OUTPUT_DIR = "pre_change_snapshots"


def load_dotenv():
    """Load KEY=VALUE pairs from a .env (or credential.env) file next to this
    script into os.environ. Real environment variables take priority."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    env_path = next((p for p in (os.path.join(script_dir, name) for name in (".env", "credential.env"))
                     if os.path.isfile(p)), None)
    if env_path is None:
        return
    with open(env_path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip().strip("'\"")
            if key and key not in os.environ:
                os.environ[key] = value


def parse_args():
    load_dotenv()
    p = argparse.ArgumentParser(description="Gather pre-change show-command snapshots from Catalyst Center devices.")
    p.add_argument("--host", default=os.environ.get("DNAC_HOST"),
                   help="Catalyst Center hostname or IP, no scheme (or set DNAC_HOST)")
    p.add_argument("--file", required=True, help="Input .xlsx or .csv file with an 'IP Address' column")
    p.add_argument("--username", default=os.environ.get("DNAC_USER"), help="API username (or set DNAC_USER)")
    p.add_argument("--password", default=os.environ.get("DNAC_PASSWORD"), help="API password (or set DNAC_PASSWORD)")
    p.add_argument("--insecure", "-k", action="store_true", help="Skip TLS certificate verification (self-signed labs)")
    args = p.parse_args()

    if not args.host:
        args.host = input("Catalyst Center host/IP: ")
    if not args.username:
        args.username = input("Catalyst Center username: ")
    if not args.password:
        args.password = getpass.getpass("Catalyst Center password: ")
    return args


def normalize_host(host):
    """Strip any scheme the caller already included (DNAC_HOST=https://... is
    a common credential.env mistake) so we don't end up with https://https://."""
    return re.sub(r"^\s*https?://", "", host.strip(), flags=re.IGNORECASE).rstrip("/")


def build_client(host, username, password, verify):
    return CatalystCenterAPI(
        base_url=f"https://{normalize_host(host)}",
        username=username,
        password=password,
        verify=verify,
        wait_on_rate_limit=True,
    )


def unwrap(response):
    """SDK responses are typically Box-like objects with a .response attribute;
    some calls (e.g. File API) return the raw content directly instead."""
    if hasattr(response, "response"):
        return response.response
    if isinstance(response, dict):
        return response.get("response", response)
    return response


def resolve_method(obj, candidates, label):
    """Find the first attribute name that actually exists on the installed
    SDK object, instead of assuming a single guessed name is correct."""
    for name in candidates:
        if hasattr(obj, name):
            return getattr(obj, name), name
    available = [m for m in dir(obj) if not m.startswith("_")]
    sys.exit(
        f"ERROR: none of {candidates} exist under api.{label} in your installed catalystcentersdk.\n"
        f"Available methods on api.{label}: {available}\n"
    )


def call_sdk(method, label, **kwargs):
    """Call an SDK method; on a parameter mismatch, print its real signature
    instead of a bare traceback, so a fix is one glance away."""
    try:
        return method(**kwargs)
    except TypeError as exc:
        sig = inspect.signature(method)
        sys.exit(
            f"ERROR: {label} rejected the parameters {list(kwargs.keys())}.\n"
            f"  {exc}\n"
            f"  Actual signature: {label}{sig}\n"
        )


DEBUG = os.environ.get("CDP_SDK_DEBUG") == "1"


def debug(msg):
    if DEBUG:
        print(f"[debug] {msg}")


# -- Step 3: IP -> device UUID (via SDK) -------------------------------------
def get_device_id(api, ip):
    method, name = resolve_method(api.devices, ["get_device_list", "get_device_by_ip", "get_network_device_by_ip"], "devices")
    response = call_sdk(method, f"devices.{name}", managementIpAddress=ip)
    result = unwrap(response)
    if isinstance(result, list):
        result = result[0] if result else None
    if not result:
        return None, None
    return result.get("id"), result.get("hostname", ip)


# -- Step 4: run one command on a batch of devices ---------------------------
def run_command(api, device_ids, command):
    method, name = resolve_method(
        api.command_runner,
        ["run_read_only_commands_on_devices_to_get_their_real_time_information",
         "run_read_only_commands_on_devices"],
        "command_runner",
    )
    response = call_sdk(method, f"command_runner.{name}", commands=[command], deviceUuids=device_ids)
    task = unwrap(response)
    task_id = task.get("taskId") if isinstance(task, dict) else getattr(task, "taskId", None)
    if not task_id:
        raise RuntimeError(f"no taskId returned from command_runner.{name}: {task}")
    return task_id


# -- Step: poll the task until it finishes -----------------------------------
def wait_for_task(api, task_id):
    method, name = resolve_method(api.task, ["get_task_by_id"], "task")
    deadline = time.time() + TASK_POLL_TIMEOUT
    while time.time() < deadline:
        response = call_sdk(method, f"task.{name}", task_id=task_id)
        task = unwrap(response)
        is_error = task.get("isError") if isinstance(task, dict) else getattr(task, "isError", False)
        end_time = task.get("endTime") if isinstance(task, dict) else getattr(task, "endTime", None)
        if is_error:
            reason = task.get("failureReason") if isinstance(task, dict) else getattr(task, "failureReason", task)
            raise RuntimeError(f"Command Runner task failed: {reason}")
        if end_time:
            progress = task.get("progress") if isinstance(task, dict) else getattr(task, "progress", "")
            try:
                file_id = json.loads(progress).get("fileId")
            except (ValueError, TypeError, AttributeError):
                file_id = None
            if not file_id:
                raise RuntimeError(f"task finished but no fileId in progress: {progress}")
            return file_id
        time.sleep(TASK_POLL_INTERVAL)
    raise RuntimeError(f"Command Runner task {task_id} did not finish within {TASK_POLL_TIMEOUT}s")


def extract_file_content(response):
    """The File API can return a raw download-wrapper object (Cisco's
    DownloadResponse) instead of already-parsed JSON."""
    if isinstance(response, (list, dict)):
        return response

    json_method = getattr(response, "json", None)
    if callable(json_method):
        try:
            return json_method()
        except Exception as exc:
            debug(f"response.json() raised: {exc}")

    for attr in ("data", "content", "text", "body", "raw_response"):
        value = getattr(response, attr, None)
        if value is None:
            continue
        if isinstance(value, (list, dict)):
            return value
        if isinstance(value, (bytes, bytearray)):
            try:
                return json.loads(value.decode("utf-8"))
            except Exception as exc:
                debug(f"response.{attr} (bytes) json.loads failed: {exc}")
        elif isinstance(value, str):
            try:
                return json.loads(value)
            except Exception as exc:
                debug(f"response.{attr} (str) json.loads failed: {exc}")

    debug(f"could not extract content from {type(response).__name__}; "
          f"available attrs: {[m for m in dir(response) if not m.startswith('_')]}")
    return []


# -- Step: fetch the raw CLI output -------------------------------------------
def get_file(api, file_id):
    method, name = resolve_method(
        api.file, ["download_a_file_by_fileid", "download_a_file_by_file_id", "get_file"], "file")
    response = call_sdk(method, f"file.{name}", file_id=file_id)
    result = unwrap(response)
    result = extract_file_content(result)
    return result


def chunked(seq, size):
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


# -- Step 2: read the input file for IP addresses -----------------------------
def read_ips(path):
    """Yield IP addresses from an .xlsx or .csv file with an 'IP Address' column."""
    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("ERROR: reading .xlsx needs openpyxl - run: pip install openpyxl "
                     "(or export the sheet as CSV instead)")
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = ([("" if c is None else str(c).strip()) for c in row]
                for row in ws.iter_rows(values_only=True))
    else:
        rows = ([c.strip() for c in row] for row in csv.reader(open(path, newline="")))

    header = next(rows, None)
    if not header:
        sys.exit(f"ERROR: {path} is empty.")
    lowered = [h.lower() for h in header]
    try:
        ip_col = lowered.index(COL_IP)
    except ValueError:
        sys.exit(f"ERROR: {path} must have an '{COL_IP}' column header (found: {header})")

    seen = set()
    for row in rows:
        if len(row) <= ip_col:
            continue
        ip = row[ip_col]
        if ip and ip not in seen:
            seen.add(ip)
            yield ip


def run_command_on_batches(api, device_ids, command):
    """Run one command across all device batches, returning
    {deviceUuid: output_text_or_None} for every device_id passed in."""
    results = {d: None for d in device_ids}
    errors = {}
    batches = list(chunked(device_ids, BATCH_SIZE))
    for i, batch in enumerate(batches, 1):
        try:
            task_id = run_command(api, batch, command)
            file_id = wait_for_task(api, task_id)
            batch_results = get_file(api, file_id)
        except Exception as exc:
            for d in batch:
                errors[d] = str(exc)
            if i < len(batches):
                time.sleep(BATCH_PAUSE)
            continue

        for entry in batch_results or []:
            entry = dict(entry) if not isinstance(entry, dict) else entry
            dev_uuid = entry.get("deviceUuid")
            command_responses = entry.get("commandResponses") or {}
            raw = (command_responses.get("SUCCESS") or {}).get(command)
            if raw:
                results[dev_uuid] = raw
            else:
                err = (command_responses.get("FAILURE") or {}).get(command) \
                    or (command_responses.get("BLACKLISTED") or {}).get(command) \
                    or "no output returned"
                errors[dev_uuid] = err
        if i < len(batches):
            time.sleep(BATCH_PAUSE)
    return results, errors


def main():
    args = parse_args()

    ips = list(read_ips(args.file))
    if not ips:
        sys.exit(f"ERROR: no usable rows found in {args.file}")
    print(f"Loaded {len(ips)} device IP(s) from {args.file}\n")

    print(f"Connecting to https://{normalize_host(args.host)} ...")
    api = build_client(args.host, args.username, args.password, verify=not args.insecure)
    print("Authenticated OK\n")

    devices = []   # (ip, device_id, hostname)
    skipped = []
    for ip in ips:
        device_id, hostname = get_device_id(api, ip)
        if device_id is None:
            print(f"!! {ip}: not found in inventory - skipping")
            skipped.append(ip)
            continue
        print(f"ok {ip}: {hostname} -> deviceId {device_id}")
        devices.append((ip, device_id, hostname))
    print()

    if not devices:
        sys.exit("ERROR: no devices resolved to inventory entries - nothing to snapshot.")

    device_ids = [d[1] for d in devices]
    id_to_source = {d[1]: (d[0], d[2]) for d in devices}

    # Create every device's file up front and write its header immediately,
    # then append each command's section as soon as that command finishes
    # for ALL devices - so a run interrupted partway through (network blip,
    # Ctrl+C, closed terminal) still leaves every file with whatever
    # commands completed before the interruption, instead of losing
    # everything (which is what happened when files were only written once
    # at the very end, after all 6 commands x all devices finished).
    script_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(script_dir, OUTPUT_DIR)
    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

    file_paths = {}   # deviceUuid -> path
    for d in device_ids:
        ip, hostname = id_to_source[d]
        safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", hostname or ip)
        out_path = os.path.join(output_dir, f"{safe_name}_{timestamp}.txt")
        with open(out_path, "w") as f:
            f.write(f"{'=' * 70}\n")
            f.write(f"Device: {hostname} ({ip})\n")
            f.write(f"Snapshot taken: {datetime.datetime.now().isoformat()}\n")
            f.write(f"{'=' * 70}\n\n")
        file_paths[d] = out_path
    print(f"Created {len(file_paths)} snapshot file(s) in {output_dir}\n")

    for cmd_idx, command in enumerate(COMMANDS, 1):
        print(f"[{cmd_idx}/{len(COMMANDS)}] Running '{command}' on {len(device_ids)} device(s)...")
        results, errors = run_command_on_batches(api, device_ids, command)
        for d in device_ids:
            output = results.get(d)
            if not output:
                output = f"[ERROR: {errors.get(d, 'no output returned')}]"
                ip, hostname = id_to_source[d]
                print(f"  !! {hostname} ({ip}): {errors.get(d, 'no output returned')}")
            with open(file_paths[d], "a") as f:
                f.write(f"##### {command} #####\n")
                f.write(output)
                f.write("\n\n")
        print()

    print("\n" + "=" * 60)
    print(f"Summary: {len(file_paths)} device snapshot(s) written to {output_dir}")
    if skipped:
        print(f"{len(skipped)} IP(s) skipped (not in inventory): {', '.join(skipped)}")


if __name__ == "__main__":
    main()
