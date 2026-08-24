#!/usr/bin/env python3
"""
Distribute (stage) the tagged Golden software image onto Catalyst Center
devices, by IP address. This ONLY distributes the image to device flash -
it never activates/applies it. Activation is a separate API call and is
deliberately not implemented here.

Reads an Excel (.xlsx) or CSV file with at least an "IP Address" column
(any other columns, e.g. Device Family / Site, are ignored).

*** Uses the legacy SWIM API, not the newer networkDeviceImages family. ***
Both /networkDeviceImages/distribute/bulk and the per-device
/networkDeviceImages/<id>/distribute returned HTTP 404 "BAPI not found"
against this deployment (confirmed: Catalyst Center 2.3.7.11) - that API
family was introduced in a later release and simply isn't registered on
this controller. This script instead uses the classic SWIM endpoints,
which have been stable since early DNA Center releases:

Workflow (all steps automated):
    1. Authenticate            POST /dna/system/api/v1/auth/token
    2. Read IP list from the input file
    3. Resolve IP -> device    GET  /dna/intent/api/v1/network-device/ip-address/<ip>
    4. Resolve Golden image    GET  /dna/intent/api/v1/image/importation?family=<family>&isTaggedGolden=true
                                per unique device family (cached - not repeated
                                per device). The classic distribute endpoint has
                                no auto-select behavior, so this step is required
                                unless --image-id is given explicitly.
    5. Trigger distribution    POST /dna/intent/api/v1/image/distribution
                                body is an ARRAY of {deviceUuid, imageUuid} pairs -
                                ALL devices are sent in one call (this endpoint's
                                schema is natively bulk, unlike the newer family)
    6. Poll the task           GET  /dna/intent/api/v1/task/<taskId>
                                the classic task API - isError/endTime/progress
    7. Print a summary of what succeeded / failed / was skipped

No activation step is performed anywhere in this script.

Usage:
    python distribute_golden_image.py --file devices.xlsx
    python distribute_golden_image.py --file devices.csv --insecure --dry-run

Credentials and host are taken from (highest priority first):
    1. command-line flags --host / --username / --password
    2. environment variables DNAC_HOST / DNAC_USER / DNAC_PASSWORD
    3. a .env (or credential.env) file in the same folder as this script
Anything still missing is prompted for interactively (password hidden).
Keep the .env file out of git.

Requires: requests. For .xlsx input also: openpyxl (CSV needs nothing extra).
"""

import argparse
import csv
import getpass
import os
import re
import sys
import time

import requests
import urllib3

COL_IP = "ip address"

DISTRIBUTE_RETRIES = 3    # extra attempts when Catalyst Center returns 429/5xx
RETRY_DELAY = 15          # seconds before first retry (doubles each attempt)
POLL_INTERVAL = 10        # seconds between task-status polls
POLL_TIMEOUT = 3600       # give up polling after this many seconds (a large batch's image copy can take a while)


def load_dotenv():
    """Load KEY=VALUE pairs from a .env (or credential.env) file next to this
    script into os.environ. Real environment variables take priority."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    env_path = next((p for p in (os.path.join(script_dir, name) for name in (".env", "credential.env"))
                     if os.path.isfile(p)), None)
    if env_path is None:
        return
    with open(env_path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip().strip("'\"")
            if key and key not in os.environ:
                os.environ[key] = value


def parse_args():
    load_dotenv()
    p = argparse.ArgumentParser(description="Distribute the Golden image to Catalyst Center devices (legacy SWIM API, no activation).")
    p.add_argument("--host", default=os.environ.get("DNAC_HOST"),
                   help="Catalyst Center hostname or IP, no scheme (or set DNAC_HOST)")
    p.add_argument("--file", required=True, help="Input .xlsx or .csv file with an 'IP Address' column")
    p.add_argument("--username", default=os.environ.get("DNAC_USER"), help="API username (or set DNAC_USER)")
    p.add_argument("--password", default=os.environ.get("DNAC_PASSWORD"), help="API password (or set DNAC_PASSWORD)")
    p.add_argument("--image-id", default=None,
                   help="Optional specific software image UUID to distribute to every device, "
                        "instead of looking up each device family's tagged Golden image")
    p.add_argument("--insecure", "-k", action="store_true", help="Skip TLS certificate verification (self-signed labs)")
    p.add_argument("--dry-run", action="store_true", help="Resolve devices and Golden images but do not trigger distribution")
    args = p.parse_args()

    if not args.host:
        args.host = input("Catalyst Center host/IP: ")
    if not args.username:
        args.username = input("Catalyst Center username: ")
    if not args.password:
        args.password = getpass.getpass("Catalyst Center password: ")
    return args


def normalize_host(host):
    """Strip any scheme the caller already included (DNAC_HOST=https://... is
    a common credential.env mistake) so we don't end up with https://https://."""
    return re.sub(r"^\s*https?://", "", host.strip(), flags=re.IGNORECASE).rstrip("/")


class DnacClient:
    def __init__(self, host, username, password, verify=True):
        self.base = f"https://{normalize_host(host)}"
        self.verify = verify
        if not verify:
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        self.session = requests.Session()
        self.session.verify = verify
        self._authenticate(username, password)

    # -- Step 1: authentication ---------------------------------------------
    def _authenticate(self, username, password):
        r = self.session.post(
            f"{self.base}/dna/system/api/v1/auth/token",
            auth=(username, password),
            timeout=30,
        )
        if r.status_code == 401:
            sys.exit("ERROR: authentication failed (401) - check username/password.")
        r.raise_for_status()
        self.session.headers.update({
            "X-Auth-Token": r.json()["Token"],
            "Content-Type": "application/json",
        })

    # -- Step 3: IP -> device record -----------------------------------------
    def get_device_by_ip(self, ip):
        """Return the inventory record for a device IP, or None if not in inventory."""
        r = self.session.get(
            f"{self.base}/dna/intent/api/v1/network-device/ip-address/{ip}",
            timeout=30,
        )
        if r.status_code in (404, 500):
            return None
        r.raise_for_status()
        return r.json().get("response") or None

    # -- Step 4: family -> Golden image UUID ---------------------------------
    def get_golden_image_id(self, family):
        """Return the imageUuid tagged Golden for this device family, or None
        if none can be safely determined.

        The device inventory's broad family ('Switches and Hubs') is not the
        same value SWIM's own catalog uses internally (e.g. 'CAT9K_LITE'), so
        the direct filtered lookup usually misses. Fall back to every image
        tagged Golden system-wide: if there's exactly one, it's unambiguous
        regardless of what its family field says, so use it. If there are
        several and none match by family, refuse to guess - list them so the
        run can be pointed at the right one via --image-id instead of
        silently picking (possibly the wrong) one.
        """
        r = self.session.get(
            f"{self.base}/dna/intent/api/v1/image/importation",
            params={"family": family, "isTaggedGolden": "true"},
            timeout=30,
        )
        r.raise_for_status()
        images = r.json().get("response") or []
        if images:
            image = images[0]
            image_id = image.get("imageUuid") or image.get("id")
            if image_id:
                return image_id
            print(f"  [debug] golden image entry has no imageUuid/id field - raw entry: {image}")

        r2 = self.session.get(
            f"{self.base}/dna/intent/api/v1/image/importation",
            params={"isTaggedGolden": "true"},
            timeout=30,
        )
        r2.raise_for_status()
        all_golden = r2.json().get("response") or []

        if not all_golden:
            print(f"  [debug] no golden-tagged images exist at all in this Catalyst Center "
                  f"- someone needs to tag an image Golden first (Design > Image Repository)")
            return None

        if len(all_golden) == 1:
            image = all_golden[0]
            image_id = image.get("imageUuid") or image.get("id")
            print(f"  [debug] exactly one Golden image exists system-wide "
                  f"(family='{image.get('family')}', version={image.get('displayVersion')}) "
                  f"- using it for family '{family}' unambiguously")
            return image_id

        print(f"  [debug] {len(all_golden)} Golden images exist and none match family='{family}' "
              f"- can't safely auto-select. Candidates:")
        for img in all_golden:
            print(f"  [debug]   imageUuid={img.get('imageUuid')} family={img.get('family')} "
                  f"version={img.get('displayVersion')}")
        print(f"  [debug] use --image-id to pin one explicitly for this run")
        return None

    # -- Step 5: trigger distribution for all devices in ONE call -----------
    # This endpoint's schema is a bare array of {deviceUuid, imageUuid} pairs,
    # unlike the newer per-device/bulk endpoints that don't exist on this
    # Catalyst Center version - so all devices go in a single request here.
    def distribute_legacy(self, pairs):
        r = self.session.post(
            f"{self.base}/dna/intent/api/v1/image/distribution",
            json=pairs,
            timeout=120,
        )
        r.raise_for_status()
        return r.json()

    # -- Step 6: poll the classic task API -----------------------------------
    def wait_for_task(self, task_id):
        start = time.time()
        deadline = start + POLL_TIMEOUT
        last_heartbeat = start
        while time.time() < deadline:
            r = self.session.get(f"{self.base}/dna/intent/api/v1/task/{task_id}", timeout=30)
            r.raise_for_status()
            task = r.json().get("response") or {}
            if task.get("isError"):
                return {"status": "FAILED", "detail": task.get("failureReason") or str(task)}
            if task.get("endTime"):
                return {"status": "SUCCESS", "detail": task.get("progress") or ""}
            if time.time() - last_heartbeat >= 60:
                elapsed = int(time.time() - start)
                print(f"  .. still waiting ({elapsed}s elapsed) - task progress so far: {task.get('progress') or 'n/a'}")
                last_heartbeat = time.time()
            time.sleep(POLL_INTERVAL)
        return {"status": "TIMEOUT", "detail": f"no result after {POLL_TIMEOUT}s"}


def http_error_detail(exc):
    """Pull the human-readable error out of a Catalyst Center error response."""
    resp = getattr(exc, "response", None)
    if resp is None:
        return str(exc)
    try:
        err = resp.json().get("errorResponse") or {}
    except ValueError:
        return f"HTTP {resp.status_code}: {resp.text[:500]}"
    parts = []
    bapi_msg = (err.get("bapiErrorResponse") or {}).get("bapiErrorMessage")
    if bapi_msg:
        parts.append(bapi_msg)
    for comp in err.get("componentErrorResponse") or []:
        msg = comp.get("componentErrorMessage") or comp.get("errorMessage")
        if msg:
            parts.append(msg)
    return "; ".join(parts) or f"HTTP {resp.status_code}: {resp.text[:500]}"


def distribute_with_retry(dnac, pairs):
    """Trigger the distribute call, retrying on transient server overload."""
    delay = RETRY_DELAY
    for attempt in range(DISTRIBUTE_RETRIES + 1):
        try:
            return dnac.distribute_legacy(pairs)
        except requests.HTTPError as exc:
            resp = exc.response
            retryable = resp is not None and resp.status_code in (429, 500, 502, 503, 504)
            if not retryable or attempt == DISTRIBUTE_RETRIES:
                raise
            print(f".. HTTP {resp.status_code} from Catalyst Center, "
                  f"waiting {delay}s then retrying ({attempt + 1}/{DISTRIBUTE_RETRIES})")
            time.sleep(delay)
            delay *= 2


# -- Step 2: read the input file for IP addresses ---------------------------
def read_ips(path):
    """Yield IP addresses from an .xlsx or .csv file with an 'IP Address' column."""
    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("ERROR: reading .xlsx needs openpyxl - run: pip install openpyxl "
                     "(or export the sheet as CSV instead)")
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = ([("" if c is None else str(c).strip()) for c in row]
                for row in ws.iter_rows(values_only=True))
    else:
        rows = ([c.strip() for c in row] for row in csv.reader(open(path, newline="")))

    header = next(rows, None)
    if not header:
        sys.exit(f"ERROR: {path} is empty.")
    lowered = [h.lower() for h in header]
    try:
        ip_col = lowered.index(COL_IP)
    except ValueError:
        sys.exit(f"ERROR: {path} must have an '{COL_IP}' column header (found: {header})")

    seen = set()
    for row in rows:
        if len(row) <= ip_col:
            continue
        ip = row[ip_col]
        if ip and ip not in seen:
            seen.add(ip)
            yield ip


def main():
    args = parse_args()

    ips = list(read_ips(args.file))
    if not ips:
        sys.exit(f"ERROR: no usable rows found in {args.file}")
    print(f"Loaded {len(ips)} device IP(s) from {args.file}\n")

    print(f"Connecting to https://{normalize_host(args.host)} ...")
    dnac = DnacClient(args.host, args.username, args.password, verify=not args.insecure)
    print("Authenticated OK\n")

    devices = []   # (ip, device_id, hostname, family)
    skipped = []
    for ip in ips:
        device = dnac.get_device_by_ip(ip)
        if device is None:
            print(f"!! {ip}: not found in inventory - skipping")
            skipped.append((ip, "device not in inventory"))
            continue
        device_id, hostname, family = device.get("id"), device.get("hostname", "?"), device.get("family", "?")
        print(f"ok {ip}: {hostname} ({family}) -> deviceId {device_id}")
        devices.append((ip, device_id, hostname, family))
    print()

    if not devices:
        sys.exit("ERROR: no devices resolved to inventory entries - nothing to distribute.")

    # Resolve each unique family's Golden image once (not per device).
    print("Resolving Golden image per device family...")
    golden_by_family = {}
    for family in sorted({d[3] for d in devices}):
        if args.image_id:
            golden_by_family[family] = args.image_id
            print(f"  {family}: using --image-id override {args.image_id}")
            continue
        image_id = dnac.get_golden_image_id(family)
        golden_by_family[family] = image_id
        print(f"  {family}: {'golden imageUuid ' + image_id if image_id else 'NO GOLDEN IMAGE FOUND'}")
    print()

    pairs = []          # [{"deviceUuid":..., "imageUuid":...}, ...]
    device_by_uuid = {}  # deviceUuid -> (ip, hostname)
    for ip, device_id, hostname, family in devices:
        image_id = golden_by_family.get(family)
        if not image_id:
            print(f"!! {ip} ({hostname}): no Golden image for family '{family}' - skipping")
            skipped.append((ip, f"no Golden image tagged for family '{family}'"))
            continue
        pairs.append({"deviceUuid": device_id, "imageUuid": image_id})
        device_by_uuid[device_id] = (ip, hostname)

    if not pairs:
        sys.exit("ERROR: no device/image pairs resolved - nothing to distribute.")

    if args.dry_run:
        print(f"[dry-run] would trigger distribution for {len(pairs)} device(s):")
        for p in pairs:
            ip, hostname = device_by_uuid[p["deviceUuid"]]
            print(f"  {ip} ({hostname}) -> image {p['imageUuid']}")
        sys.exit(0)

    print(f"Triggering distribution for {len(pairs)} device(s) in one request...")
    try:
        trigger = distribute_with_retry(dnac, pairs)
    except requests.HTTPError as exc:
        sys.exit(f"ERROR: distribute request failed: {http_error_detail(exc)}")
    print(f"  raw trigger response: {trigger}")

    task_id = (trigger.get("response") or {}).get("taskId") or trigger.get("taskId")
    if not task_id:
        sys.exit(f"ERROR: no taskId returned - raw response: {trigger}")
    print(f"\ntaskId: {task_id} - polling until the batch finishes "
          f"(image copies happen together; this can take a while)...\n")

    result = dnac.wait_for_task(task_id)
    status = result["status"]
    detail = result["detail"]

    # This classic API returns one task for the whole batch, not clean
    # per-device granularity - so every device in this run shares the same
    # outcome unless/until we see the real response and can refine this.
    ok, failed = [], []
    if status == "SUCCESS":
        print(f"=> batch distribution {status.lower()}: {detail}")
        ok = list(device_by_uuid.values())
    else:
        print(f"!! batch distribution {status.lower()}: {detail}")
        failed = [(ip, hostname, detail) for ip, hostname in device_by_uuid.values()]

    # -- Step 7: summary -------------------------------------------------------
    print("\n" + "=" * 60)
    print(f"Summary: {len(ok)} distributed, {len(failed)} failed, {len(skipped)} skipped")
    for ip, hostname, reason in failed:
        print(f"  FAILED  {ip} ({hostname}): {reason}")
    for ip, reason in skipped:
        print(f"  SKIPPED {ip}: {reason}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
