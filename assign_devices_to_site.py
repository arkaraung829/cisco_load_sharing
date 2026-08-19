#!/usr/bin/env python3
"""
Assign network devices to sites in Cisco Catalyst Center (DNA Center).

Reads an Excel (.xlsx) or CSV file with the columns:

    IP Address | Device Family | Site

where "Site" is the full site name hierarchy, e.g.
    Global/Canada/Saskatchewan/LCL Stores/LCL Stores SK1/Store 07579

Workflow (all steps automated):
    1. Authenticate           POST /dna/system/api/v1/auth/token
    2. Read + group the input file rows by site
    3. Resolve site name->ID  GET  /dna/intent/api/v1/site?name=<site>
    4. Verify device exists   GET  /dna/intent/api/v1/network-device/ip-address/<ip>
    5. Assign devices         POST /dna/intent/api/v1/assign-device-to-site/<siteId>/device
    6. Poll execution status until the assignment completes
    7. Print a summary of what succeeded / failed / was skipped

Usage:
    python assign_devices_to_site.py --host dnac.example.com --file devices.xlsx
    python assign_devices_to_site.py --host 10.1.1.10 --file devices.csv --insecure --dry-run

Credentials and host are taken from (highest priority first):
    1. command-line flags --host / --username / --password
    2. environment variables DNAC_HOST / DNAC_USER / DNAC_PASSWORD
    3. a .env file in the same folder as this script, e.g.:
           DNAC_HOST=10.1.1.10
           DNAC_USER=admin
           DNAC_PASSWORD=YourPassword123
Anything still missing is prompted for interactively (password hidden).
Keep the .env file out of git - it is listed in .gitignore.

Requires: requests. For .xlsx input also: openpyxl (CSV needs nothing extra).
"""

import argparse
import csv
import getpass
import os
import re
import sys
import time
from collections import OrderedDict

import requests
import urllib3

# Column headers expected in the input file (case-insensitive match)
COL_IP = "ip address"
COL_SITE = "site"

EXEC_POLL_INTERVAL = 2   # seconds between execution-status polls
EXEC_POLL_TIMEOUT = 120  # give up polling after this many seconds
ASSIGN_RETRIES = 3       # extra attempts when Catalyst Center returns 429/5xx
RETRY_DELAY = 15         # seconds before first retry (doubles each attempt)
PAUSE_BETWEEN_SITES = 2  # seconds between per-site assignment calls


def load_dotenv():
    """Load KEY=VALUE pairs from a .env (or credential.env) file next to this
    script into os.environ.

    Real environment variables take priority over .env values. Lines starting
    with '#' and blank lines are ignored; optional surrounding quotes stripped.
    """
    script_dir = os.path.dirname(os.path.abspath(__file__))
    env_path = next((p for p in (os.path.join(script_dir, name) for name in (".env", "credential.env"))
                     if os.path.isfile(p)), None)
    if env_path is None:
        return
    with open(env_path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip().strip("'\"")
            if key and key not in os.environ:
                os.environ[key] = value


def parse_args():
    load_dotenv()
    p = argparse.ArgumentParser(description="Assign devices to Catalyst Center sites from an Excel/CSV file.")
    p.add_argument("--host", default=os.environ.get("DNAC_HOST"),
                   help="Catalyst Center hostname or IP, no scheme (or set DNAC_HOST)")
    p.add_argument("--file", required=True, help="Input .xlsx or .csv file with IP Address / Site columns")
    p.add_argument("--username", default=os.environ.get("DNAC_USER"), help="API username (or set DNAC_USER)")
    p.add_argument("--password", default=os.environ.get("DNAC_PASSWORD"), help="API password (or set DNAC_PASSWORD)")
    p.add_argument("--insecure", "-k", action="store_true", help="Skip TLS certificate verification (self-signed labs)")
    p.add_argument("--dry-run", action="store_true", help="Resolve sites and devices but do not assign anything")
    args = p.parse_args()

    if not args.host:
        args.host = input("Catalyst Center host/IP: ")
    if not args.username:
        args.username = input("Catalyst Center username: ")
    if not args.password:
        args.password = getpass.getpass("Catalyst Center password: ")
    return args


def normalize_host(host):
    """Strip any scheme the caller already included (DNAC_HOST=https://... is
    a common credential.env mistake) so we don't end up with https://https://."""
    return re.sub(r"^\s*https?://", "", host.strip(), flags=re.IGNORECASE).rstrip("/")


class DnacClient:
    def __init__(self, host, username, password, verify=True):
        self.base = f"https://{normalize_host(host)}"
        self.verify = verify
        if not verify:
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        self.session = requests.Session()
        self.session.verify = verify
        self._authenticate(username, password)

    # -- Step 1: authentication -------------------------------------------
    def _authenticate(self, username, password):
        r = self.session.post(
            f"{self.base}/dna/system/api/v1/auth/token",
            auth=(username, password),
            timeout=30,
        )
        if r.status_code == 401:
            sys.exit("ERROR: authentication failed (401) - check username/password.")
        r.raise_for_status()
        self.session.headers.update({
            "X-Auth-Token": r.json()["Token"],
            "Content-Type": "application/json",
        })

    # -- Step 3: site name -> site ID -------------------------------------
    def get_site_id(self, site_name):
        """Return the site UUID for a full hierarchy name, or None if not found."""
        r = self.session.get(
            f"{self.base}/dna/intent/api/v1/site",
            params={"name": site_name},
            timeout=30,
        )
        if r.status_code == 404:
            return None
        r.raise_for_status()
        response = r.json().get("response") or []
        if not response:
            return None
        return response[0]["id"]

    # -- Step 4: device inventory check -----------------------------------
    def get_device_by_ip(self, ip):
        """Return the inventory record for a device IP, or None if not in inventory."""
        r = self.session.get(
            f"{self.base}/dna/intent/api/v1/network-device/ip-address/{ip}",
            timeout=30,
        )
        if r.status_code in (404, 500):
            # Catalyst Center returns an error payload when the IP is unknown
            return None
        r.raise_for_status()
        return r.json().get("response") or None

    # -- Step 5: assign devices to a site ---------------------------------
    def assign_devices(self, site_id, ips):
        body = {"device": [{"ip": ip} for ip in ips]}
        r = self.session.post(
            f"{self.base}/dna/intent/api/v1/assign-device-to-site/{site_id}/device",
            json=body,
            headers={"__runsync": "true", "__persistbapioutput": "true"},
            timeout=60,
        )
        r.raise_for_status()
        return r.json()

    # -- Step 6: poll execution status until done -------------------------
    def wait_for_execution(self, result):
        """Follow executionStatusUrl until the BAPI execution finishes."""
        status_url = result.get("executionStatusUrl")
        if not status_url:
            return result  # synchronous response already contains the outcome
        deadline = time.time() + EXEC_POLL_TIMEOUT
        while time.time() < deadline:
            r = self.session.get(f"{self.base}{status_url}", timeout=30)
            r.raise_for_status()
            status = r.json()
            if status.get("status") not in (None, "IN_PROGRESS", "PENDING"):
                return status
            time.sleep(EXEC_POLL_INTERVAL)
        return {"status": "TIMEOUT", "bapiError": f"no result after {EXEC_POLL_TIMEOUT}s"}


def http_error_detail(exc):
    """Pull the human-readable error out of a Catalyst Center error response."""
    resp = getattr(exc, "response", None)
    if resp is None:
        return str(exc)
    try:
        err = resp.json().get("errorResponse") or {}
    except ValueError:
        return f"HTTP {resp.status_code}: {resp.text[:500]}"
    parts = []
    bapi_msg = (err.get("bapiErrorResponse") or {}).get("bapiErrorMessage")
    if bapi_msg:
        parts.append(bapi_msg)
    for comp in err.get("componentErrorResponse") or []:
        msg = comp.get("componentErrorMessage") or comp.get("errorMessage")
        if msg:
            parts.append(msg)
    return "; ".join(parts) or f"HTTP {resp.status_code}: {resp.text[:500]}"


def assign_with_retry(dnac, site_id, ips):
    """Assign, retrying with growing waits when the server is overloaded."""
    delay = RETRY_DELAY
    for attempt in range(ASSIGN_RETRIES + 1):
        try:
            result = dnac.assign_devices(site_id, ips)
            return dnac.wait_for_execution(result)
        except requests.HTTPError as exc:
            resp = exc.response
            retryable = resp is not None and resp.status_code in (429, 500, 502, 503, 504)
            if not retryable or attempt == ASSIGN_RETRIES:
                raise
            print(f"  .. HTTP {resp.status_code} from Catalyst Center, "
                  f"waiting {delay}s then retrying ({attempt + 1}/{ASSIGN_RETRIES})")
            time.sleep(delay)
            delay *= 2


# -- Step 2: read the input file and group rows by site -------------------
def read_rows(path):
    """Yield (ip, site) tuples from an .xlsx or .csv file."""
    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("ERROR: reading .xlsx needs openpyxl - run: pip install openpyxl "
                     "(or export the sheet as CSV instead)")
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = ([("" if c is None else str(c).strip()) for c in row]
                for row in ws.iter_rows(values_only=True))
    else:
        rows = ([c.strip() for c in row] for row in csv.reader(open(path, newline="")))

    header = next(rows, None)
    if not header:
        sys.exit(f"ERROR: {path} is empty.")
    lowered = [h.lower() for h in header]
    try:
        ip_col = lowered.index(COL_IP)
        site_col = lowered.index(COL_SITE)
    except ValueError:
        sys.exit(f"ERROR: {path} must have '{COL_IP}' and '{COL_SITE}' column headers "
                 f"(found: {header})")

    for row in rows:
        if len(row) <= max(ip_col, site_col):
            continue
        ip, site = row[ip_col], row[site_col]
        if ip and site:
            yield ip, site


def group_by_site(rows):
    """Group device IPs by site, preserving file order and dropping duplicates."""
    grouped = OrderedDict()
    for ip, site in rows:
        grouped.setdefault(site, [])
        if ip not in grouped[site]:
            grouped[site].append(ip)
    return grouped


def main():
    args = parse_args()

    grouped = group_by_site(read_rows(args.file))
    if not grouped:
        sys.exit(f"ERROR: no usable rows found in {args.file}")
    total_devices = sum(len(v) for v in grouped.values())
    print(f"Loaded {total_devices} device(s) across {len(grouped)} site(s) from {args.file}\n")

    print(f"Connecting to https://{normalize_host(args.host)} ...")
    dnac = DnacClient(args.host, args.username, args.password, verify=not args.insecure)
    print("Authenticated OK\n")

    ok, failed, skipped = [], [], []

    for site_name, ips in grouped.items():
        print(f"=== {site_name} ===")

        site_id = dnac.get_site_id(site_name)
        if not site_id:
            print(f"  !! site not found in Catalyst Center - skipping {len(ips)} device(s)\n")
            skipped.extend((ip, site_name, "site not found") for ip in ips)
            continue
        print(f"  siteId: {site_id}")

        # Only send devices that actually exist in inventory
        assignable = []
        for ip in ips:
            device = dnac.get_device_by_ip(ip)
            if device is None:
                print(f"  !! {ip}: not found in inventory - skipping")
                skipped.append((ip, site_name, "device not in inventory"))
            else:
                print(f"  ok {ip}: {device.get('hostname', '?')} ({device.get('family', '?')})")
                assignable.append(ip)

        if not assignable:
            print()
            continue

        if args.dry_run:
            print(f"  [dry-run] would assign {len(assignable)} device(s) to this site\n")
            ok.extend((ip, site_name) for ip in assignable)
            continue

        try:
            status = assign_with_retry(dnac, site_id, assignable)
        except requests.HTTPError as exc:
            detail = http_error_detail(exc)
            print(f"  !! assignment request failed: {detail}\n")
            failed.extend((ip, site_name, detail) for ip in assignable)
            continue
        finally:
            time.sleep(PAUSE_BETWEEN_SITES)

        # Depending on the Catalyst Center version, a successful synchronous
        # run reports status "SUCCESS" or the string "True" (with the human-
        # readable outcome in result.progress).
        outcome = str(status.get("status", "")).upper()
        progress = (status.get("result") or {}).get("progress", "")
        if not status.get("bapiError") and outcome in ("SUCCESS", "TRUE", ""):
            print(f"  => {progress or status.get('message') or 'assigned'}\n")
            ok.extend((ip, site_name) for ip in assignable)
        else:
            err = status.get("bapiError") or progress or status.get("message") or str(status)
            print(f"  !! assignment failed: {err}\n")
            failed.extend((ip, site_name, err) for ip in assignable)

    # -- Step 7: summary ---------------------------------------------------
    print("=" * 60)
    label = "would be assigned (dry-run)" if args.dry_run else "assigned"
    print(f"Summary: {len(ok)} {label}, {len(failed)} failed, {len(skipped)} skipped")
    for ip, site, reason in failed:
        print(f"  FAILED  {ip} -> {site}: {reason}")
    for ip, site, reason in skipped:
        print(f"  SKIPPED {ip} -> {site}: {reason}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
