#!/usr/bin/env python3
"""
Activate the tagged Golden software image on Catalyst Center devices, by IP
address. The image must already be present on the device (distributed
first - see distribute_golden_image.py) unless --distribute-if-needed is
given.

*** THIS IS A LIVE, SERVICE-IMPACTING OPERATION. ***
Activation applies the staged image, which on Catalyst switches means a
REBOOT - this is not a background copy like distribution. Running this
for real against a device causes a real outage window for whatever it
serves, for as long as the reboot/boot-up takes.

*** --upgrade-mode has no default and must be supplied explicitly. ***
The documented request schema lists 'deviceUpgradeMode' as a plain string
with no enumerated valid values shown to this script's author - guessing
a value for a field that controls how a live device boots is not
acceptable. Confirm the correct value against Cisco's schema doc
(https://developer.cisco.com/docs/catalyst-center/triggersoftwareimageactivationrequest)
before running anything but --dry-run.

Uses the legacy SWIM API - same family already confirmed working on this
Catalyst Center (2.3.7.11) by distribute_golden_image.py, since the newer
networkDeviceImages family returned "BAPI not found" on this deployment.

Workflow:
    1. Authenticate            POST /dna/system/api/v1/auth/token
    2. Read IP list from the input file
    3. Resolve IP -> device    GET  /dna/intent/api/v1/network-device/ip-address/<ip>
    4. Resolve Golden image    GET  /dna/intent/api/v1/image/importation?family=<family>&isTaggedGolden=true
                                same lookup/fallback logic as distribute_golden_image.py
    5. Trigger activation      POST /dna/intent/api/v1/image/activation/device
                                body is an array of activation request objects, exactly
                                matching the documented schema - no extra fields are sent
    6. Poll the task           GET  /dna/intent/api/v1/task/<taskId>
    7. Print a summary of what succeeded / failed / was skipped

Deliberately NOT implemented, per instruction:
    - No task-name field - not in the documented request schema (same
      reasoning applied to distribute_golden_image.py: undocumented
      fields are not sent to this production system).
    - No scheduled/deferred activation - out of scope for this version.
      scheduleValidate (query param) only validates a schedule that lives
      elsewhere in the broader workflow; it is not sent by this script.

Usage:
    python activate_golden_image.py --file devices.csv --upgrade-mode <value> --dry-run
    python activate_golden_image.py --file devices.csv --upgrade-mode <value> --insecure

A real (non-dry-run) run requires typed confirmation unless --yes is
given, since this can reboot production devices.

Credentials and host are taken from (highest priority first):
    1. command-line flags --host / --username / --password
    2. environment variables DNAC_HOST / DNAC_USER / DNAC_PASSWORD
    3. a .env (or credential.env) file in the same folder as this script
Anything still missing is prompted for interactively (password hidden).

Requires: requests. For .xlsx input also: openpyxl (CSV needs nothing extra).
"""

import argparse
import csv
import getpass
import os
import re
import sys
import time

import requests
import urllib3

COL_IP = "ip address"

ACTIVATE_RETRIES = 3      # extra attempts when Catalyst Center returns 429/5xx
RETRY_DELAY = 15          # seconds before first retry (doubles each attempt)
POLL_INTERVAL = 10        # seconds between task-status polls
POLL_TIMEOUT = 1800       # give up polling after this many seconds (reboot + boot-up takes a while)


def load_dotenv():
    """Load KEY=VALUE pairs from a .env (or credential.env) file next to this
    script into os.environ. Real environment variables take priority."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    env_path = next((p for p in (os.path.join(script_dir, name) for name in (".env", "credential.env"))
                     if os.path.isfile(p)), None)
    if env_path is None:
        return
    with open(env_path) as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, _, value = line.partition("=")
            key, value = key.strip(), value.strip().strip("'\"")
            if key and key not in os.environ:
                os.environ[key] = value


def parse_args():
    load_dotenv()
    p = argparse.ArgumentParser(
        description="Activate the Golden image on Catalyst Center devices (legacy SWIM API). "
                    "REBOOTS the device - review before running without --dry-run.")
    p.add_argument("--host", default=os.environ.get("DNAC_HOST"),
                   help="Catalyst Center hostname or IP, no scheme (or set DNAC_HOST)")
    p.add_argument("--file", required=True, help="Input .xlsx or .csv file with an 'IP Address' column")
    p.add_argument("--username", default=os.environ.get("DNAC_USER"), help="API username (or set DNAC_USER)")
    p.add_argument("--password", default=os.environ.get("DNAC_PASSWORD"), help="API password (or set DNAC_PASSWORD)")
    p.add_argument("--upgrade-mode", required=True,
                   help="REQUIRED, no default. Value for the documented 'deviceUpgradeMode' field - "
                        "confirm the correct value against Cisco's schema doc before using this for real.")
    p.add_argument("--image-id", default=None,
                   help="Optional specific software image UUID to activate on every device, "
                        "instead of looking up each device family's tagged Golden image")
    p.add_argument("--activate-lower-version", action="store_true",
                   help="Sets activateLowerImageVersion=true (allow activating an image older "
                        "than what's currently running). Default false.")
    p.add_argument("--distribute-if-needed", action="store_true",
                   help="Sets distributeIfNeeded=true (let Catalyst Center distribute the image "
                        "first if it isn't already on the device). Default false - distribute "
                        "explicitly first with distribute_golden_image.py instead.")
    p.add_argument("--insecure", "-k", action="store_true", help="Skip TLS certificate verification (self-signed labs)")
    p.add_argument("--dry-run", action="store_true", help="Resolve devices and Golden images but do not trigger activation")
    p.add_argument("--yes", action="store_true",
                   help="Skip the typed confirmation prompt before a real (non-dry-run) activation run")
    args = p.parse_args()

    if not args.host:
        args.host = input("Catalyst Center host/IP: ")
    if not args.username:
        args.username = input("Catalyst Center username: ")
    if not args.password:
        args.password = getpass.getpass("Catalyst Center password: ")
    return args


def normalize_host(host):
    """Strip any scheme the caller already included (DNAC_HOST=https://... is
    a common credential.env mistake) so we don't end up with https://https://."""
    return re.sub(r"^\s*https?://", "", host.strip(), flags=re.IGNORECASE).rstrip("/")


class DnacClient:
    def __init__(self, host, username, password, verify=True):
        self.base = f"https://{normalize_host(host)}"
        self.verify = verify
        if not verify:
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        self.session = requests.Session()
        self.session.verify = verify
        self._authenticate(username, password)

    # -- Step 1: authentication ---------------------------------------------
    def _authenticate(self, username, password):
        r = self.session.post(
            f"{self.base}/dna/system/api/v1/auth/token",
            auth=(username, password),
            timeout=30,
        )
        if r.status_code == 401:
            sys.exit("ERROR: authentication failed (401) - check username/password.")
        r.raise_for_status()
        self.session.headers.update({
            "X-Auth-Token": r.json()["Token"],
            "Content-Type": "application/json",
        })

    # -- Step 3: IP -> device record -----------------------------------------
    def get_device_by_ip(self, ip):
        """Return the inventory record for a device IP, or None if not in inventory."""
        r = self.session.get(
            f"{self.base}/dna/intent/api/v1/network-device/ip-address/{ip}",
            timeout=30,
        )
        if r.status_code in (404, 500):
            return None
        r.raise_for_status()
        return r.json().get("response") or None

    # -- Step 4: family -> Golden image UUID ---------------------------------
    # Identical logic to distribute_golden_image.py's DnacClient.get_golden_image_id -
    # kept in sync deliberately since both scripts need the same resolution.
    def get_golden_image_id(self, family):
        """Return the imageUuid tagged Golden for this device family, or None
        if none can be safely determined."""
        r = self.session.get(
            f"{self.base}/dna/intent/api/v1/image/importation",
            params={"family": family, "isTaggedGolden": "true"},
            timeout=30,
        )
        r.raise_for_status()
        images = r.json().get("response") or []
        if images:
            image = images[0]
            image_id = image.get("imageUuid") or image.get("id")
            if image_id:
                return image_id
            print(f"  [debug] golden image entry has no imageUuid/id field - raw entry: {image}")

        r2 = self.session.get(
            f"{self.base}/dna/intent/api/v1/image/importation",
            params={"isTaggedGolden": "true"},
            timeout=30,
        )
        r2.raise_for_status()
        all_golden = r2.json().get("response") or []

        if not all_golden:
            print(f"  [debug] no golden-tagged images exist at all in this Catalyst Center "
                  f"- someone needs to tag an image Golden first (Design > Image Repository)")
            return None

        if len(all_golden) == 1:
            image = all_golden[0]
            image_id = image.get("imageUuid") or image.get("id")
            print(f"  [debug] exactly one Golden image exists system-wide "
                  f"(family='{image.get('family')}', version={image.get('displayVersion')}) "
                  f"- using it for family '{family}' unambiguously")
            return image_id

        print(f"  [debug] {len(all_golden)} Golden images exist and none match family='{family}' "
              f"- can't safely auto-select. Candidates:")
        for img in all_golden:
            print(f"  [debug]   imageUuid={img.get('imageUuid')} family={img.get('family')} "
                  f"version={img.get('displayVersion')}")
        print(f"  [debug] use --image-id to pin one explicitly for this run")
        return None

    # -- Step 5: trigger activation for all devices in ONE call -------------
    # Body matches the documented schema exactly - no extra/undocumented
    # fields (e.g. no task name) are ever sent.
    def activate_legacy(self, entries):
        r = self.session.post(
            f"{self.base}/dna/intent/api/v1/image/activation/device",
            json=entries,
            timeout=120,
        )
        r.raise_for_status()
        return r.json()

    # -- Step 6: poll the classic task API -----------------------------------
    def wait_for_task(self, task_id):
        start = time.time()
        deadline = start + POLL_TIMEOUT
        last_heartbeat = start
        while time.time() < deadline:
            r = self.session.get(f"{self.base}/dna/intent/api/v1/task/{task_id}", timeout=30)
            r.raise_for_status()
            task = r.json().get("response") or {}
            if task.get("isError"):
                return {"status": "FAILED", "detail": task.get("failureReason") or str(task)}
            if task.get("endTime"):
                return {"status": "SUCCESS", "detail": task.get("progress") or ""}
            if time.time() - last_heartbeat >= 60:
                elapsed = int(time.time() - start)
                print(f"  .. still waiting ({elapsed}s elapsed) - task progress so far: {task.get('progress') or 'n/a'}")
                last_heartbeat = time.time()
            time.sleep(POLL_INTERVAL)
        return {"status": "TIMEOUT", "detail": f"no result after {POLL_TIMEOUT}s"}


def http_error_detail(exc):
    """Pull the human-readable error out of a Catalyst Center error response."""
    resp = getattr(exc, "response", None)
    if resp is None:
        return str(exc)
    try:
        err = resp.json().get("errorResponse") or {}
    except ValueError:
        return f"HTTP {resp.status_code}: {resp.text[:500]}"
    parts = []
    bapi_msg = (err.get("bapiErrorResponse") or {}).get("bapiErrorMessage")
    if bapi_msg:
        parts.append(bapi_msg)
    for comp in err.get("componentErrorResponse") or []:
        msg = comp.get("componentErrorMessage") or comp.get("errorMessage")
        if msg:
            parts.append(msg)
    return "; ".join(parts) or f"HTTP {resp.status_code}: {resp.text[:500]}"


def activate_with_retry(dnac, entries):
    """Trigger the activate call, retrying on transient server overload."""
    delay = RETRY_DELAY
    for attempt in range(ACTIVATE_RETRIES + 1):
        try:
            return dnac.activate_legacy(entries)
        except requests.HTTPError as exc:
            resp = exc.response
            retryable = resp is not None and resp.status_code in (429, 500, 502, 503, 504)
            if not retryable or attempt == ACTIVATE_RETRIES:
                raise
            print(f".. HTTP {resp.status_code} from Catalyst Center, "
                  f"waiting {delay}s then retrying ({attempt + 1}/{ACTIVATE_RETRIES})")
            time.sleep(delay)
            delay *= 2


# -- Step 2: read the input file for IP addresses ---------------------------
def read_ips(path):
    """Yield IP addresses from an .xlsx or .csv file with an 'IP Address' column."""
    if path.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("ERROR: reading .xlsx needs openpyxl - run: pip install openpyxl "
                     "(or export the sheet as CSV instead)")
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = ([("" if c is None else str(c).strip()) for c in row]
                for row in ws.iter_rows(values_only=True))
    else:
        rows = ([c.strip() for c in row] for row in csv.reader(open(path, newline="")))

    header = next(rows, None)
    if not header:
        sys.exit(f"ERROR: {path} is empty.")
    lowered = [h.lower() for h in header]
    try:
        ip_col = lowered.index(COL_IP)
    except ValueError:
        sys.exit(f"ERROR: {path} must have an '{COL_IP}' column header (found: {header})")

    seen = set()
    for row in rows:
        if len(row) <= ip_col:
            continue
        ip = row[ip_col]
        if ip and ip not in seen:
            seen.add(ip)
            yield ip


def main():
    args = parse_args()

    ips = list(read_ips(args.file))
    if not ips:
        sys.exit(f"ERROR: no usable rows found in {args.file}")
    print(f"Loaded {len(ips)} device IP(s) from {args.file}\n")

    print(f"Connecting to https://{normalize_host(args.host)} ...")
    dnac = DnacClient(args.host, args.username, args.password, verify=not args.insecure)
    print("Authenticated OK\n")

    devices = []   # (ip, device_id, hostname, family)
    skipped = []
    for ip in ips:
        device = dnac.get_device_by_ip(ip)
        if device is None:
            print(f"!! {ip}: not found in inventory - skipping")
            skipped.append((ip, "device not in inventory"))
            continue
        device_id, hostname, family = device.get("id"), device.get("hostname", "?"), device.get("family", "?")
        print(f"ok {ip}: {hostname} ({family}) -> deviceId {device_id}")
        devices.append((ip, device_id, hostname, family))
    print()

    if not devices:
        sys.exit("ERROR: no devices resolved to inventory entries - nothing to activate.")

    print("Resolving Golden image per device family...")
    golden_by_family = {}
    for family in sorted({d[3] for d in devices}):
        if args.image_id:
            golden_by_family[family] = args.image_id
            print(f"  {family}: using --image-id override {args.image_id}")
            continue
        image_id = dnac.get_golden_image_id(family)
        golden_by_family[family] = image_id
        print(f"  {family}: {'golden imageUuid ' + image_id if image_id else 'NO GOLDEN IMAGE FOUND'}")
    print()

    entries = []          # matches triggerSoftwareImageActivationRequest[] exactly
    device_by_uuid = {}   # deviceUuid -> (ip, hostname)
    for ip, device_id, hostname, family in devices:
        image_id = golden_by_family.get(family)
        if not image_id:
            print(f"!! {ip} ({hostname}): no Golden image for family '{family}' - skipping")
            skipped.append((ip, f"no Golden image tagged for family '{family}'"))
            continue
        entries.append({
            "deviceUpgradeMode": args.upgrade_mode,
            "deviceUuid": device_id,
            "activateLowerImageVersion": args.activate_lower_version,
            "distributeIfNeeded": args.distribute_if_needed,
            "imageUuidList": [image_id],
            "smuImageUuidList": [],
        })
        device_by_uuid[device_id] = (ip, hostname)

    if not entries:
        sys.exit("ERROR: no device/image entries resolved - nothing to activate.")

    print(f"[plan] would activate {len(entries)} device(s), upgradeMode='{args.upgrade_mode}', "
          f"activateLowerImageVersion={args.activate_lower_version}, "
          f"distributeIfNeeded={args.distribute_if_needed}:")
    for e in entries:
        ip, hostname = device_by_uuid[e["deviceUuid"]]
        print(f"  {ip} ({hostname}) -> image {e['imageUuidList'][0]}")
    print()

    if args.dry_run:
        print("[dry-run] stopping here - no activation triggered.")
        sys.exit(0)

    print("*" * 70)
    print("WARNING: activation applies the image and typically REBOOTS the")
    print("device(s) listed above. This is a live, service-impacting change.")
    print("*" * 70)
    if not args.yes:
        confirm = input(f"\nType ACTIVATE to proceed with {len(entries)} device(s), anything else to abort: ")
        if confirm.strip() != "ACTIVATE":
            sys.exit("Aborted - no activation triggered.")

    print(f"\nTriggering activation for {len(entries)} device(s) in one request...")
    try:
        trigger = activate_with_retry(dnac, entries)
    except requests.HTTPError as exc:
        sys.exit(f"ERROR: activate request failed: {http_error_detail(exc)}")
    print(f"  raw trigger response: {trigger}")

    task_id = (trigger.get("response") or {}).get("taskId") or trigger.get("taskId")
    if not task_id:
        sys.exit(f"ERROR: no taskId returned - raw response: {trigger}")
    print(f"\ntaskId: {task_id} - polling until the batch finishes "
          f"(reboot + boot-up can take several minutes per device)...\n")

    result = dnac.wait_for_task(task_id)
    status = result["status"]
    detail = result["detail"]

    # This classic API returns one task for the whole batch, not clean
    # per-device granularity - same caveat as distribute_golden_image.py.
    ok, failed = [], []
    if status == "SUCCESS":
        print(f"=> batch activation {status.lower()}: {detail}")
        ok = list(device_by_uuid.values())
    else:
        print(f"!! batch activation {status.lower()}: {detail}")
        failed = [(ip, hostname, detail) for ip, hostname in device_by_uuid.values()]

    # -- Step 7: summary -------------------------------------------------------
    print("\n" + "=" * 60)
    print(f"Summary: {len(ok)} activated, {len(failed)} failed, {len(skipped)} skipped")
    for ip, hostname, reason in failed:
        print(f"  FAILED  {ip} ({hostname}): {reason}")
    for ip, reason in skipped:
        print(f"  SKIPPED {ip}: {reason}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
